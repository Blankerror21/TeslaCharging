"""One-time migration from the original Electric_Usage.xlsx into data/readings.csv.

Kept for provenance: it documents exactly how the spreadsheet was interpreted,
including the two date typos that were corrected. Requires openpyxl; nothing
else in this repo does.

Usage: python3 scripts/migrate_from_xlsx.py Electric_Usage.xlsx
"""

import csv
import re
import sys
from datetime import date

import openpyxl

# Row ranges of the two stacked blocks in the original sheet.
USAGE_ROWS = range(3, 13)  # "Electric usage": 120v in col B, 240v in col C
WALL_ROWS = range(18, 36)  # "Tesla Wall Connector": kWh in col C

RANGE_RE = re.compile(r"^\s*(\d{1,2}/\d{1,2}/\d{2,4})\s*-\s*(\d{1,2}/\d{1,2}/\d{2,4})\s*$")


def parse_day(text):
    month, day, year = (int(part) for part in text.split("/"))
    return date(2000 + year if year < 100 else year, month, day)


def parse_range(text):
    match = RANGE_RE.match(text)
    if not match:
        raise ValueError(f"unparseable date range: {text!r}")
    start, end = parse_day(match.group(1)), parse_day(match.group(2))
    if end < start:
        # Rows 11 and 25 both read "12/x/24 - 1/4/24"; the end year is a typo.
        end = end.replace(year=end.year + 1)
    return start, end


def main(path):
    sheet = openpyxl.load_workbook(path, data_only=True)["Sheet1"]
    rows = []

    for row in USAGE_ROWS:
        label = sheet.cell(row, 1).value
        if not label:
            continue
        start, end = parse_range(label)
        for column, meter in ((2, "meter_120v"), (3, "meter_240v")):
            kwh = sheet.cell(row, column).value
            if kwh is not None:
                rows.append((start, end, meter, float(kwh), ""))

    for row in WALL_ROWS:
        label = sheet.cell(row, 1).value
        if not label:
            continue
        kwh = sheet.cell(row, 3).value
        if kwh is None:
            # Row 35 (10/1/25 - 10/30/25) has no reading entered.
            print(f"skipping row {row} ({label.strip()}): no kWh recorded", file=sys.stderr)
            continue
        start, end = parse_range(label)
        rows.append((start, end, "wall_connector", float(kwh), ""))

    rows.sort(key=lambda r: (r[0], r[2]))

    writer = csv.writer(sys.stdout, lineterminator="\n")
    writer.writerow(["start", "end", "meter", "kwh", "note"])
    for start, end, meter, kwh, note in rows:
        writer.writerow([start.isoformat(), end.isoformat(), meter, f"{kwh:g}", note])


if __name__ == "__main__":
    main(sys.argv[1])
